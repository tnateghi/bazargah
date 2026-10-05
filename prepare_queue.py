from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from db import init_db, replace_queue

DEFAULT_ACTIVITY = "دام سبك و سنگین روستایی عشایری"
ACTIVITY_VALUE = "1114"  # دام سبك و سنگین روستایی عشایری


def _norm(x) -> str | None:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    s = str(x).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if not s or s.lower() == "nan":
        return None
    return s


def _tonum(x) -> int:
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return 0
        if isinstance(x, str) and not x.strip():
            return 0
        return int(float(x))
    except Exception:
        return 0


def load_stats_xlsx(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["سراب"]
    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if all(v is None or (isinstance(v, str) and not str(v).strip()) for v in vals):
            continue
        rows.append(vals)
    wb.close()

    headers = [(h if h is not None else f"col{i}") for i, h in enumerate(rows[0])]
    maxc = max(len(r) for r in rows)
    while len(headers) < maxc:
        headers.append(f"col{len(headers)}")
    data = [list(r) + [None] * (maxc - len(r)) for r in rows[1:]]
    df = pd.DataFrame(data, columns=headers[:maxc])

    sheep_m = next(c for c in df.columns if "گوسفند" in str(c) and "نر" in str(c))
    sheep_f = next(c for c in df.columns if "گوسفند" in str(c) and "ماده" in str(c))
    goat_m = next(
        c for c in df.columns if str(c).replace("\n", "").startswith("بز") and "نر" in str(c)
    )
    goat_f = next(
        c
        for c in df.columns
        if str(c).replace("\n", "").startswith("بز") and "ماده" in str(c)
    )
    heavy = next(c for c in df.columns if "جمع کل دام سنگین" in str(c))

    out = pd.DataFrame(
        {
            "shenase": df["شناسه یکتا"].map(_norm),
            "name_stats": df["دامدار"].astype(str),
            "light_count": (
                df[sheep_m].map(_tonum)
                + df[sheep_f].map(_tonum)
                + df[goat_m].map(_tonum)
                + df[goat_f].map(_tonum)
            ),
            "heavy_count": df[heavy].map(_tonum),
        }
    )
    return out.dropna(subset=["shenase"]).drop_duplicates("shenase")


def load_site_report(path: Path) -> pd.DataFrame:
    df = pd.read_html(path)[0]
    # ستون «ردیف» گزارش سایت = کد تفصیلی (مثلاً 1054576)، نه شماره ترتیبی لیست
    radif_col = next((c for c in df.columns if str(c).strip() == "ردیف"), None)
    out = pd.DataFrame(
        {
            "sort_order": range(1, len(df) + 1),
            "tafsili_code": df[radif_col].map(_norm) if radif_col else None,
            "shenase": df["شناسه یکتا"].map(_norm),
            "name": df["نام زیرمجموعه"].astype(str),
            "old_light": df["تعداد دام سبک"].map(_tonum),
            "old_heavy": df["تعداد دام سنگین"].map(_tonum),
            "activity_type": df["نوع فعالیت"].fillna(DEFAULT_ACTIVITY).astype(str),
        }
    )
    return out.dropna(subset=["shenase"]).drop_duplicates("shenase")


def build_queue(stats_path: Path, report_path: Path, activity: str | None = None) -> dict:
    stats = load_stats_xlsx(stats_path)
    report = load_site_report(report_path)
    # keep report order
    merged = report.merge(stats, on="shenase", how="inner").sort_values("sort_order")
    if activity:
        merged["activity_type"] = activity

    rows = []
    for r in merged.itertuples(index=False):
        changed = int(
            not (
                int(r.old_light) == int(r.light_count)
                and int(r.old_heavy) == int(r.heavy_count)
            )
        )
        # if already same numbers on site, no need to resend
        status = "nochange" if changed == 0 else "pending"
        rows.append(
            {
                "sort_order": int(r.sort_order),
                "tafsili_code": getattr(r, "tafsili_code", None),
                "shenase": r.shenase,
                "name": r.name,
                "light_count": int(r.light_count),
                "heavy_count": int(r.heavy_count),
                "activity_type": r.activity_type or DEFAULT_ACTIVITY,
                "old_light": int(r.old_light),
                "old_heavy": int(r.old_heavy),
                "changed": changed,
                "status": status,
            }
        )
    init_db()
    return replace_queue(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build work queue from two Excel files")
    parser.add_argument("--stats", required=True, help="آمار xlsx path")
    parser.add_argument("--report", required=True, help="ZirMajmoeReport xls/html path")
    parser.add_argument("--activity", default=None, help="Override activity type for all rows")
    args = parser.parse_args()
    s = build_queue(Path(args.stats), Path(args.report), args.activity)
    print(
        f"Queue ready. total={s['total']} pending={s['pending']} "
        f"done={s['done']} changed={s['changed']} unchanged={s['unchanged']}"
    )


if __name__ == "__main__":
    main()
